"""
ALTECH SOFTWARE DEVELOPERS
SMART DOCUMENT INTELLIGENCE ENGINE v26
--------------------------------------
Universal document ingestion via a unified document_engine package.
Unstructured is used as an optional helper for complex files across
ALL supported types (Excel, CSV, DOCX, PDF, images, text).

Endpoints (unchanged from v25):
  GET  /                     service info
  GET  /api/ping             liveness
  GET  /api/health           full status
  POST /api/analyze          extract items from a file
  POST /api/chat             natural-language command execution
  POST /api/apply-prompt     apply a prompt to already-extracted items
  POST /api/lookup           free web lookup
  POST /api/extract-text     raw text of a file
  POST /api/match-products   fuzzy product matching
"""

from __future__ import annotations

import io
import os
import re
import json
import math
import time
import logging
import urllib.request
import urllib.parse
import urllib.error
from pathlib import Path
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
import openpyxl
import PyPDF2
import docx
from PIL import Image, ImageEnhance, ImageOps
import pytesseract
from rapidfuzz import fuzz

from fastapi import FastAPI, File, UploadFile, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None

try:
    from chat_engine import process_message as chat_process_message
    _CHAT_AVAILABLE = True
except Exception as _e:
    logging.getLogger("altech-engine").warning(
        f"chat_engine not available: {_e}")
    chat_process_message = None
    _CHAT_AVAILABLE = False

# ---------------------------------------------------------------------------
# NEW: Unified document engine (Unstructured-aware, universal file support)
# ---------------------------------------------------------------------------
try:
    from document_engine.pipeline import analyze_document
    _DOC_ENGINE_AVAILABLE = True
except Exception as _e:
    logging.getLogger("altech-engine").warning(
        f"document_engine not available: {_e}")
    analyze_document = None
    _DOC_ENGINE_AVAILABLE = False


logging.basicConfig(level=logging.INFO)
log = logging.getLogger("altech-smart-document")

ENGINE_VERSION = "26.0.0"
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "40"))
MAX_ROWS = int(os.getenv("MAX_ROWS", "20000"))
MAX_SHEETS = int(os.getenv("MAX_SHEETS", "50"))
MAX_OCR_PAGES = int(os.getenv("MAX_OCR_PAGES", "30"))
OCR_DPI = int(os.getenv("OCR_DPI", "200"))
WEB_LOOKUP_ENABLED = os.getenv("WEB_LOOKUP_ENABLED", "1") == "1"
WEB_LOOKUP_TIMEOUT = int(os.getenv("WEB_LOOKUP_TIMEOUT", "6"))

TESS_CMD = os.getenv("TESSERACT_CMD", "/usr/bin/tesseract")
if os.path.exists(TESS_CMD):
    pytesseract.pytesseract.tesseract_cmd = TESS_CMD


app = FastAPI(title="Altech Smart Document Intelligence Engine",
              version=ENGINE_VERSION)
app.add_middleware(CORSMiddleware, allow_origins=["*"],
                   allow_methods=["*"], allow_headers=["*"])


# ===========================================================================
#  ROLE MODEL
# ===========================================================================
ROLE_ALIASES = {
    "product_name": [
        "product", "product name", "item", "item name", "flower",
        "flower name", "description", "item description",
        "particulars", "goods", "article", "articles",
        "flower variety", "variety", "cultivar",
    ],
    "variety": ["variety", "cultivar"],
    "quantity": [
        "quantity", "qty", "qnty", "stems", "total stems", "pieces",
        "pcs", "units", "count", "total quantity",
    ],
    "boxes": ["boxes", "box", "bx", "cartons", "carton", "ctn", "cases"],
    "pack_rate": [
        "packrate", "pack rate", "pack_rate", "per box", "per carton",
        "stems per box", "stems/box", "qty per box", "quantity per box",
    ],
    "unit_price": [
        "price", "unit price", "unit_price", "rate", "cost",
        "unit cost", "price per stem", "price/stem",
    ],
    "total": [
        "total", "amount", "line total", "line amount",
        "total amount", "extended price", "revenue",
    ],
    "length_cm": [
        "length", "length cm", "length (cm)", "length(cm)",
        "stem length", "size",
    ],
    "head_size_cm": ["head size", "head size cm", "head size (cm)"],
    "color": ["color", "colour", "shade"],
    "farm_code": ["farm code", "farm", "farmcode"],
    "invoice_number": ["invoice number", "invoice no", "invoice #",
                       "reference"],
    "date": ["date", "invoice date", "shipment date"],
    "due_date": ["due date", "payment due", "valid until"],
    "currency": ["currency"],
    "vat_rate": ["vat", "vat rate", "tax", "tax rate"],
    "discount": ["discount"],
    "n": ["n", "no", "no.", "#", "s/n", "sr", "index"],
}

NUMERIC_ROLES = {
    "quantity", "boxes", "pack_rate", "length_cm", "head_size_cm",
    "unit_price", "total", "vat_rate", "discount", "n",
}

ROLE_PRIORITY = {
    "quantity": ["total stems", "number of stems", "stem quantity",
                 "stems", "total quantity", "quantity", "qty"],
    "unit_price": ["unit price", "price per stem", "unit cost",
                   "rate", "price", "cost"],
    "total": ["line total", "total amount", "extended price",
              "line amount", "amount", "total"],
    "boxes": ["number of boxes", "no. of boxes", "cartons", "cases",
              "boxes", "box"],
}


def norm(s: Any) -> str:
    s = "" if s is None else str(s)
    s = s.lower().replace("–", "-").replace("—", "-").replace("’", "'")
    return re.sub(r"\s+", " ", s.strip())


def clean_key(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", norm(s)).strip("_")


# ===========================================================================
#  NUMBER PARSING
# ===========================================================================
def parse_number(v: Any) -> Optional[float]:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return x if math.isfinite(x) else None
    s = str(v).strip()
    if not s:
        return None
    if s.startswith("="):
        return None
    s = re.sub(
        r"(?i)\b(?:usd|us\$|kes|ksh|kshs|eur|gbp|aed|sar|qar|"
        r"dollars?|shillings?)\b", "", s)
    s = s.replace("$", "").replace("€", "").replace("£", "").strip()
    if "," in s and "." in s:
        if s.rfind(".") > s.rfind(","):
            s = s.replace(",", "")
        else:
            s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        if re.search(r",\d{1,2}$", s):
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
    s = s.replace("\u00a0", "").replace("\u202f", "")
    s = re.sub(r"(?<=\d)\s+(?=\d)", "", s)
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return None
    try:
        return float(m.group())
    except ValueError:
        return None


# ===========================================================================
#  ROLE CLASSIFICATION
# ===========================================================================
def _classify_basic(label: Any) -> Optional[str]:
    l = re.sub(r"[^a-z0-9 ]+", " ", norm(label)).strip()
    if not l:
        return None
    best_role, best_score = None, 0
    for role, names in ROLE_ALIASES.items():
        for n in names:
            a = re.sub(r"[^a-z0-9 ]+", " ", norm(n)).strip()
            if l == a:
                return role
            s = max(fuzz.WRatio(l, a), fuzz.token_set_ratio(l, a))
            if s > best_score:
                best_score, best_role = s, role
    return best_role if best_score >= 78 else None


def build_columns(headers: List[Any]) -> List[Dict[str, Any]]:
    taken: set = set()
    cols = []
    for h in headers:
        label = str(h or "").strip() or "Column"
        key = clean_key(label) or "column"
        k = key
        i = 2
        while k in taken:
            k = f"{key}_{i}"
            i += 1
        taken.add(k)
        cols.append({
            "key": k, "label": label,
            "role": _classify_basic(label),
            "source": "original",
        })

    role_groups: Dict[str, List[Dict[str, Any]]] = {}
    for c in cols:
        if c["role"]:
            role_groups.setdefault(c["role"], []).append(c)

    for role, group in role_groups.items():
        if len(group) <= 1:
            continue
        priority = ROLE_PRIORITY.get(role, [])
        winner = None
        for pref in priority:
            pref_n = re.sub(r"[^a-z0-9 ]+", " ", norm(pref)).strip()
            for c in group:
                if re.sub(r"[^a-z0-9 ]+", " ",
                          norm(c["label"])).strip() == pref_n:
                    winner = c
                    break
            if winner:
                break
        if not winner:
            winner = group[0]
        for c in group:
            if c is not winner:
                c["role"] = f"{role}__{c['key']}"

    return cols


# ===========================================================================
#  CELL / ROW CLEANING
# ===========================================================================
def clean_cell(v: Any) -> Any:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except Exception:
        pass
    if isinstance(v, (pd.Timestamp, datetime)):
        return v.isoformat()
    if isinstance(v, (int, float)):
        x = float(v)
        return int(x) if x.is_integer() else x
    return str(v).strip()


def row_is_empty(row: Dict[str, Any]) -> bool:
    return all(v in (None, "") for v in row.values())


def row_is_summary(row: Dict[str, Any]) -> bool:
    values = [v for v in row.values() if v not in (None, "")]
    if not values:
        return False
    joined = " ".join(norm(v) for v in values)
    if re.search(
        r"\b(?:subtotal|grand\s+total|invoice\s+total|total\s+amount|"
        r"balance\s+due|thank\s+you|amount\s+due|total\s+price)\b",
        joined):
        return True
    numeric_like = 0
    for v in values:
        if isinstance(v, (int, float)):
            numeric_like += 1
        elif isinstance(v, str) and re.fullmatch(
                r"\s*[\$€£]?\s*-?\d[\d,.\s]*\s*", v):
            numeric_like += 1
    if len(values) <= 3 and numeric_like == len(values):
        return True
    return False


# ===========================================================================
#  HEADER DETECTION
# ===========================================================================
def find_header_row(df: pd.DataFrame) -> int:
    best = (0, 0)
    scan = min(len(df), 40)
    for i in range(scan):
        try:
            vals = [str(x).strip() for x in df.iloc[i].tolist()
                    if pd.notna(x) and str(x).strip()]
        except Exception:
            continue
        if len(vals) < 2:
            continue
        recognized = sum(1 for x in vals if _classify_basic(x))
        if recognized == 0:
            continue
        textish = sum(1 for x in vals if re.search(r"[A-Za-z]", x))
        score = recognized * 10 + min(textish, 10) - i
        if score > best[0]:
            best = (score, i)
    return best[1]


# ===========================================================================
#  DATAFRAME → STRUCTURE
# ===========================================================================
def dataframe_to_structure(df: pd.DataFrame,
                           sheet_name: str = "") -> Dict[str, Any]:
    if df is None or df.empty:
        return {"columns": [], "items": [], "sheet": sheet_name,
                "header_row": None}

    df = df.iloc[:MAX_ROWS, :]
    header_idx = find_header_row(df)

    try:
        raw_headers = [clean_cell(x) for x in df.iloc[header_idx].tolist()]
    except Exception:
        return {"columns": [], "items": [], "sheet": sheet_name,
                "header_row": None}

    while raw_headers and raw_headers[-1] in (None, ""):
        raw_headers.pop()
    if not raw_headers:
        return {"columns": [], "items": [], "sheet": sheet_name,
                "header_row": None}

    columns = build_columns(raw_headers)
    items: List[Dict[str, Any]] = []

    for ridx in range(header_idx + 1, len(df)):
        try:
            row_vals = list(df.iloc[ridx].tolist())
        except Exception:
            continue
        row_vals = row_vals[:len(columns)]
        record: Dict[str, Any] = {}
        for j, col in enumerate(columns):
            raw = row_vals[j] if j < len(row_vals) else None
            cell = clean_cell(raw)
            if col.get("role") in NUMERIC_ROLES:
                num = parse_number(cell)
                record[col["key"]] = num if num is not None else cell
            else:
                record[col["key"]] = cell

        if row_is_empty(record):
            continue
        if row_is_summary(record):
            continue

        has_text = False
        for col in columns:
            if col.get("role") in ("product_name", "variety", "description",
                                    "color"):
                v = record.get(col["key"])
                if isinstance(v, str) and len(v.strip()) >= 2 \
                        and re.search(r"[A-Za-z]", v):
                    has_text = True
                    break
        if not has_text:
            for col in columns:
                v = record.get(col["key"])
                if isinstance(v, str) and len(v.strip()) >= 2 \
                        and re.search(r"[A-Za-z]", v):
                    has_text = True
                    break
        if not has_text:
            continue

        for col in columns:
            role = col.get("role")
            if role and role not in record:
                record[role] = record.get(col["key"])
            elif role and record.get(role) in (None, ""):
                record[role] = record.get(col["key"])

        items.append(record)

    return {
        "columns": columns,
        "items": items,
        "sheet": sheet_name,
        "header_row": header_idx + 1,
    }


# ===========================================================================
#  EXCEL / CSV
# ===========================================================================
def extract_excel(content: bytes, ext: str) -> List[Dict[str, Any]]:
    if ext == "csv":
        for sep in (",", ";", "\t"):
            try:
                df = pd.read_csv(io.BytesIO(content), header=None,
                                 sep=sep, engine="python", dtype=object)
                if df.shape[1] > 1 or sep == ",":
                    return [dataframe_to_structure(df, "CSV")]
            except Exception:
                continue
        return []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content),
                                     read_only=True, data_only=True)
    except Exception as e:
        log.warning(f"openpyxl failed: {e}")
        return []

    structures = []
    for ws in list(wb.worksheets)[:MAX_SHEETS]:
        try:
            rows = list(ws.values)
        except Exception:
            continue
        if not rows:
            continue
        df = pd.DataFrame(rows)
        structure = dataframe_to_structure(df, ws.title)
        if structure.get("items") or structure.get("columns"):
            structures.append(structure)
    return structures


# ===========================================================================
#  OCR
# ===========================================================================
def ocr_image(img: Image.Image) -> str:
    img = ImageOps.exif_transpose(img).convert("L")
    w, h = img.size
    if max(w, h) < 2200:
        s = 2200 / max(w, h)
        img = img.resize((int(w * s), int(h * s)))
    img = ImageEnhance.Contrast(img).enhance(1.7)
    best = ""
    for psm in (6, 4, 11):
        try:
            t = pytesseract.image_to_string(
                img, lang="eng", config=f"--oem 3 --psm {psm}", timeout=25)
            if len(re.findall(r"[A-Za-z0-9]", t)) > \
               len(re.findall(r"[A-Za-z0-9]", best)):
                best = t
        except Exception:
            continue
    return best


# ===========================================================================
#  PDF EXTRACTION — layout-aware
# ===========================================================================
def extract_pdf(content: bytes) -> Tuple[str, str]:
    if fitz is None:
        try:
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            text = "\n".join(p.extract_text() or "" for p in reader.pages)
            return text, "pdf_text_pypdf2"
        except Exception:
            return "", "pdf_no_reader"

    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception:
        return "", "pdf_open_failed"

    pages: List[Dict[str, Any]] = []
    for page_no, page in enumerate(doc):
        try:
            plain = page.get_text("text", sort=True) or ""
        except Exception:
            plain = ""

        blocks: List[Dict[str, Any]] = []
        try:
            for b in page.get_text("blocks", sort=True) or []:
                x0, y0, x1, y1 = b[0], b[1], b[2], b[3]
                text = (b[4] or "").strip()
                if text:
                    blocks.append({
                        "x0": float(x0), "y0": float(y0),
                        "x1": float(x1), "y1": float(y1),
                        "text": text,
                    })
        except Exception:
            pass

        pages.append({
            "page_no": page_no,
            "text": plain,
            "blocks": blocks,
            "char_count": len(re.sub(r"\s+", "", plain)),
        })

    if not pages:
        return "", "pdf_empty"

    merged = _merge_pdf_pages(pages)
    if merged:
        return merged, "pdf_merged"

    combined = "\n".join(p["text"] for p in pages if p["text"])
    if len(re.sub(r"\s+", "", combined)) >= 30:
        return combined, "pdf_text"

    try:
        parts = []
        for i, p in enumerate(doc):
            if i >= MAX_OCR_PAGES:
                break
            pix = p.get_pixmap(dpi=OCR_DPI, alpha=False)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            parts.append(ocr_image(img))
        return "\n".join(parts), "pdf_ocr"
    except Exception:
        return combined, "pdf_text_partial"


def _blocks_to_rows(blocks: List[Dict[str, Any]],
                    y_tol: float = 4.0) -> List[List[Dict[str, Any]]]:
    if not blocks:
        return []
    sorted_blocks = sorted(blocks, key=lambda b: (b["y0"], b["x0"]))
    rows: List[List[Dict[str, Any]]] = []
    current: List[Dict[str, Any]] = [sorted_blocks[0]]
    current_y = sorted_blocks[0]["y0"]
    for b in sorted_blocks[1:]:
        if abs(b["y0"] - current_y) <= y_tol:
            current.append(b)
        else:
            rows.append(sorted(current, key=lambda x: x["x0"]))
            current = [b]
            current_y = b["y0"]
    if current:
        rows.append(sorted(current, key=lambda x: x["x0"]))
    return rows


def _count_columns(blocks: List[Dict[str, Any]]) -> int:
    buckets = set()
    for b in blocks:
        cx = (b["x0"] + b["x1"]) / 2
        buckets.add(int(cx // 25))
    return len(buckets)


def _merge_pdf_pages(pages: List[Dict[str, Any]]) -> str:
    if not pages:
        return ""

    for p in pages:
        p["column_count"] = _count_columns(p["blocks"])
        p["rows"] = _blocks_to_rows(p["blocks"])

    master = max(pages, key=lambda p: p["column_count"])
    if master["column_count"] < 3:
        return ""

    indexed: Dict[int, List[str]] = {}
    orphans: List[List[str]] = []

    for p in sorted(pages, key=lambda p: p["page_no"]):
        for row_blocks in p["rows"]:
            cells = [b["text"] for b in row_blocks]
            if not cells:
                continue
            joined = " ".join(cells)
            if len(cells) == 1 and not re.search(r"[A-Za-z0-9]", joined):
                continue

            idx = None
            for j, c in enumerate(cells):
                n = parse_number(c)
                if n is not None and float(n).is_integer() \
                        and 1 <= n <= 99999:
                    idx = int(n)
                    cells_no_idx = cells[:j] + cells[j+1:]
                    break
            else:
                cells_no_idx = cells

            if idx is None:
                orphans.append(cells_no_idx)
                continue

            if idx in indexed:
                existing = indexed[idx]
                for c in cells_no_idx:
                    if c and c not in existing:
                        existing.append(c)
            else:
                indexed[idx] = list(cells_no_idx)

    if not indexed and not orphans:
        return ""

    out_lines: List[str] = []
    for idx in sorted(indexed.keys()):
        cells = indexed[idx]
        out_lines.append(" | ".join([str(idx)] + cells))
    for o in orphans:
        if o:
            out_lines.append(" | ".join(o))

    if not out_lines:
        return ""

    result = "\n".join(out_lines)
    if len(out_lines) < 3 or not re.search(r"[A-Za-z]", result):
        return ""
    return result


# ===========================================================================
#  DOCX
# ===========================================================================
def extract_docx(content: bytes) -> str:
    try:
        d = docx.Document(io.BytesIO(content))
        parts = []
        for p in d.paragraphs:
            if p.text.strip():
                parts.append(p.text)
        for table in d.tables:
            for row in table.rows:
                vals = [c.text.strip() for c in row.cells]
                if any(vals):
                    parts.append(" | ".join(vals))
        return "\n".join(parts)
    except Exception:
        return ""


# ===========================================================================
#  TEXT → STRUCTURE
# ===========================================================================
def text_to_structure(text: str) -> Dict[str, Any]:
    lines = [x.strip() for x in str(text or "").splitlines() if x.strip()]
    if not lines:
        return {"columns": [], "items": [], "sections": []}

    best_score, best_idx, best_cells = 0, None, None
    for i, line in enumerate(lines[:80]):
        cells = [x.strip() for x in re.split(r"\s*\|\s*|\t+|\s{2,}", line)
                 if x.strip()]
        if len(cells) < 2:
            continue
        score = sum(1 for c in cells if _classify_basic(c))
        if "|" in line:
            score += 2
        if score > best_score:
            best_score, best_idx, best_cells = score, i, cells

    if best_idx is None or best_score == 0:
        return {
            "columns": [{"key": "raw_text", "label": "Text",
                         "role": None, "source": "original"}],
            "items": [{"raw_text": x, "product_name": x} for x in lines],
            "sections": [],
        }

    columns = build_columns(best_cells)
    items: List[Dict[str, Any]] = []

    for line in lines[best_idx + 1:]:
        low = line.lower()
        if re.search(r"\b(?:subtotal|grand total|balance due|"
                     r"thank you|invoice total)\b", low):
            continue
        cells = [x.strip() for x in re.split(r"\s*\|\s*|\t+|\s{2,}", line)
                 if x.strip()]
        if len(cells) < 2:
            continue
        record: Dict[str, Any] = {}
        for j, col in enumerate(columns):
            raw = cells[j] if j < len(cells) else None
            if col.get("role") in NUMERIC_ROLES:
                num = parse_number(raw)
                record[col["key"]] = num if num is not None else raw
            else:
                record[col["key"]] = raw
            if col.get("role"):
                record[col["role"]] = record[col["key"]]
        if row_is_empty(record):
            continue

        has_text = False
        for col in columns:
            if col.get("role") in ("product_name", "variety", "description"):
                v = record.get(col["key"])
                if isinstance(v, str) and len(v.strip()) >= 2 \
                        and re.search(r"[A-Za-z]", v):
                    has_text = True
                    break
        if not has_text:
            for col in columns:
                v = record.get(col["key"])
                if isinstance(v, str) and len(v.strip()) >= 2 \
                        and re.search(r"[A-Za-z]", v):
                    has_text = True
                    break
        if not has_text:
            continue
        items.append(record)

    return {"columns": columns, "items": items, "sections": []}


# ===========================================================================
#  CANONICALIZATION
# ===========================================================================
def canonicalize(structure: Dict[str, Any]) -> List[Dict[str, Any]]:
    cols = structure.get("columns", [])
    items = structure.get("items", [])

    def role_col(role):
        for c in cols:
            if c.get("role") == role:
                return c

    out = []
    for i, row in enumerate(items):
        x = dict(row)
        warnings = []
        confidence = 0.55

        q_col = role_col("quantity")
        p_col = role_col("unit_price")
        t_col = role_col("total")
        b_col = role_col("boxes")
        pr_col = role_col("pack_rate")

        qv = parse_number(x.get(q_col["key"])) if q_col else None
        pv = parse_number(x.get(p_col["key"])) if p_col else None
        tv = parse_number(x.get(t_col["key"])) if t_col else None
        bv = parse_number(x.get(b_col["key"])) if b_col else None
        prv = parse_number(x.get(pr_col["key"])) if pr_col else None

        if qv is not None:
            confidence += 0.08
        if pv is not None:
            confidence += 0.08
        if tv is not None:
            confidence += 0.08

        if qv is not None and pv is not None:
            expected = qv * pv
            if t_col is None:
                x["total"] = round(expected, 4)
                confidence += 0.06
            elif tv is None:
                x[t_col["key"]] = round(expected, 4)
                confidence += 0.06
            elif abs(expected - tv) > max(0.05, abs(tv) * 0.02):
                warnings.append("quantity_x_unit_price_mismatch")

        if bv is not None and prv is not None and qv is not None:
            if abs(bv * prv - qv) > max(0.5, abs(qv) * 0.02):
                warnings.append("boxes_x_packrate_mismatch")

        x["_meta"] = {
            "row_index": i + 1,
            "confidence": round(min(confidence, 1), 3),
            "warnings": warnings,
        }
        out.append(x)
    return out


def infer_document_type(text: str, items: List[Dict[str, Any]],
                        filename: str) -> str:
    blob = norm(" ".join([Path(filename).stem, (text or "")[:5000]]))
    for kw, dt in [
        ("proforma", "proforma_invoice"),
        ("purchase order", "purchase_order"),
        ("packing list", "packing_list"),
        ("delivery note", "delivery_note"),
        ("quotation", "quotation"),
        ("quote", "quotation"),
        ("receipt", "receipt"),
        ("invoice", "invoice"),
    ]:
        if kw in blob:
            return dt
    return "line_item_document" if items else "business_document"


# ===========================================================================
#  LEGACY ANALYSIS (fallback if document_engine is missing)
# ===========================================================================
def analyze_bytes(content: bytes, fname: str, ext: str,
                  company_id: int = 0, prompt: str = "") -> Dict[str, Any]:
    started = time.perf_counter()
    ext = (ext or Path(fname).suffix.lstrip(".")).lower()
    structures: List[Dict[str, Any]] = []
    raw_text = ""
    method = "unknown"

    if ext in ("xlsx", "xlsm", "xls"):
        structures = extract_excel(content, ext)
        method = "excel"
    elif ext == "csv":
        structures = extract_excel(content, "csv")
        method = "csv"
    elif ext == "pdf":
        raw_text, method = extract_pdf(content)
        if raw_text:
            structures = [text_to_structure(raw_text)]
    elif ext in ("docx", "doc"):
        raw_text = extract_docx(content)
        method = "docx"
        if raw_text:
            structures = [text_to_structure(raw_text)]
    elif ext in ("png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"):
        try:
            raw_text = ocr_image(Image.open(io.BytesIO(content)))
            method = "image_ocr"
        except Exception:
            raw_text = ""
        if raw_text:
            structures = [text_to_structure(raw_text)]
    elif ext == "json":
        try:
            obj = json.loads(content.decode("utf-8", errors="ignore"))
            raw_text = json.dumps(obj, ensure_ascii=False, indent=2)
            if isinstance(obj, list) and obj and isinstance(obj[0], dict):
                headers = list(obj[0].keys())
                structures = [{
                    "columns": build_columns(headers),
                    "items": obj,
                    "sheet": "JSON",
                }]
            else:
                structures = [text_to_structure(raw_text)]
            method = "json"
        except Exception:
            raw_text = content.decode("utf-8", errors="ignore")
            structures = [text_to_structure(raw_text)]
            method = "text"
    else:
        raw_text = content.decode("utf-8", errors="ignore")
        structures = [text_to_structure(raw_text)]
        method = "text"

    if not structures:
        return {
            "success": False,
            "engine_version": ENGINE_VERSION,
            "error": "No readable structure was detected.",
            "columns": [], "items": [], "raw_items": [],
            "raw_text": raw_text[:20000],
            "diagnostics": {"stage": "ingestion", "file": fname,
                            "method": method, "trust_score": 0,
                            "confidence_band": "review"},
        }

    primary_columns = structures[0].get("columns", [])
    merged_items = []
    sheet_names = []
    for s in structures:
        sheet_names.append(s.get("sheet") or "")
        for row in s.get("items", []):
            r = dict(row)
            r["_sheet"] = s.get("sheet") or None
            merged_items.append(r)

    items = canonicalize({"columns": primary_columns,
                          "items": merged_items})

    total_qty = sum(parse_number(r.get("quantity")) or 0 for r in items)
    total_amount = sum(parse_number(r.get("total")) or 0 for r in items)

    trust = (sum(r["_meta"]["confidence"] for r in items) / len(items)
             if items else 0.25)
    anomalies = [
        {"row": r["_meta"]["row_index"], "code": w}
        for r in items for w in r["_meta"]["warnings"]
    ]
    trust = round(max(0, min(1, trust - 0.03 * len(anomalies))), 3)

    return {
        "success": True,
        "engine_version": ENGINE_VERSION,
        "filename": fname,
        "file_type": ext,
        "extraction_method": method,
        "document_type": infer_document_type(raw_text, items, fname),
        "columns": primary_columns,
        "items": items,
        "raw_items": items,
        "item_count": len(items),
        "total_quantity": int(total_qty) if total_qty else 0,
        "total_amount": round(total_amount, 2),
        "raw_text": raw_text[:20000],
        "diagnostics": {
            "sheets": sheet_names,
            "trust_score": trust,
            "confidence_band": ("high" if trust >= 0.85
                                else "medium" if trust >= 0.65
                                else "review"),
            "anomalies": anomalies,
        },
        "elapsed_seconds": round(time.perf_counter() - started, 3),
        "company_id": company_id,
        "prompt_received": bool(prompt),
    }


# ===========================================================================
#  WEB LOOKUP# ===========================================================================
def web_lookup(query: str) -> Dict[str, Any]:
    if not WEB_LOOKUP_ENABLED or not query:
        return {"ok": False, "reason": "disabled or empty"}
    try:
        url = ("https://api.duckduckgo.com/?q="
               + urllib.parse.quote(query)
               + "&format=json&no_html=1&skip_disambig=1")
        req = urllib.request.Request(
            url, headers={"User-Agent": "AltechSmartDocs/26.0"})
        with urllib.request.urlopen(req, timeout=WEB_LOOKUP_TIMEOUT) as r:
            data = json.loads(r.read().decode("utf-8", errors="ignore"))
        for key in ("AbstractText", "Answer", "Definition"):
            v = data.get(key)
            if v and isinstance(v, str) and len(v.strip()) > 4:
                return {"ok": True, "answer": v.strip()[:600],
                        "source": "duckduckgo"}
        for topic in data.get("RelatedTopics") or []:
            if isinstance(topic, dict) and topic.get("Text"):
                return {"ok": True, "answer": topic["Text"].strip()[:600],
                        "source": "duckduckgo"}
    except Exception as e:
        return {"ok": False, "reason": str(e)}
    return {"ok": False, "reason": "no answer found"}


# ===========================================================================
#  MODELS
# ===========================================================================
class ChatRequest(BaseModel):
    items: List[Dict[str, Any]] = []
    columns: List[Dict[str, Any]] = []
    message: str = ""
    history: List[Any] = []


class PromptRequest(BaseModel):
    items: List[Dict[str, Any]] = []
    columns: List[Dict[str, Any]] = []
    prompt: str = ""


class MatchRequest(BaseModel):
    items: List[Dict[str, Any]] = []
    company_products: List[Dict[str, Any]] = []


class LookupRequest(BaseModel):
    query: str = ""


# ===========================================================================
#  ENDPOINTS
# ===========================================================================
@app.get("/")
async def root():
    return {
        "service": "Altech Smart Document Intelligence Engine",
        "version": ENGINE_VERSION,
        "status": "operational",
        "chat_engine_available": _CHAT_AVAILABLE,
        "document_engine_available": _DOC_ENGINE_AVAILABLE,
        "web_lookup_available": WEB_LOOKUP_ENABLED,
    }


@app.get("/api/ping")
def ping():
    return {"ok": True, "version": ENGINE_VERSION,
            "chat_engine_available": _CHAT_AVAILABLE,
            "document_engine_available": _DOC_ENGINE_AVAILABLE,
            "web_lookup_available": WEB_LOOKUP_ENABLED}


@app.get("/api/health")
def health():
    return {"status": "healthy", "version": ENGINE_VERSION,
            "ocr_available": True, "pdf_available": fitz is not None,
            "chat_engine_available": _CHAT_AVAILABLE,
            "document_engine_available": _DOC_ENGINE_AVAILABLE,
            "web_lookup_available": WEB_LOOKUP_ENABLED}


@app.post("/api/analyze")
async def analyze(
    file: UploadFile = File(...),
    company_id: int = Form(0),
    file_type: Optional[str] = Form(None),
    prompt: str = Form(""),
):
    content = await file.read()
    if not content:
        raise HTTPException(400, "Empty file")
    if len(content) > MAX_UPLOAD_MB * 1024 * 1024:
        raise HTTPException(413, f"File larger than {MAX_UPLOAD_MB} MB")
    fname = file.filename or "upload"
    ext = (file_type or Path(fname).suffix.lstrip(".")).lower()

    # Preferred: unified document_engine pipeline (Unstructured-aware)
    if _DOC_ENGINE_AVAILABLE and analyze_document is not None:
        try:
            return analyze_document(content, fname, ext, company_id)
        except Exception as e:
            log.exception("document_engine failed; falling back to legacy")

    # Legacy fallback
    try:
        return analyze_bytes(content, fname, ext, company_id, prompt)
    except Exception as e:
        log.exception("Analysis failed")
        return {"success": False, "engine_version": ENGINE_VERSION,
                "error": "Analysis failed safely.",
                "columns": [], "items": [], "raw_items": [],
                "diagnostics": {"exception": str(e), "file": fname}}


@app.post("/api/chat")
async def chat(req: ChatRequest):
    if not _CHAT_AVAILABLE or chat_process_message is None:
        return {"success": False, "status": "error",
                "items": req.items, "columns": req.columns,
                "explanation": "Command engine is not loaded.",
                "needs_clarification": False}
    try:
        try:
            return chat_process_message(req.items, req.columns,
                                        req.message, req.history)
        except TypeError:
            return chat_process_message(req.items, req.columns,
                                        req.message)
    except Exception as e:
        log.exception("chat failure")
        return {"success": False, "status": "error",
                "items": req.items, "columns": req.columns,
                "explanation": "Command could not be applied safely.",
                "error": str(e), "needs_clarification": False}


@app.post("/api/apply-prompt")
async def apply_prompt(req: PromptRequest):
    return await chat(ChatRequest(items=req.items, columns=req.columns,
                                   message=req.prompt))


@app.post("/api/lookup")
async def lookup(req: LookupRequest):
    return web_lookup((req.query or "").strip())


@app.post("/api/extract-text")
async def extract_text(file: UploadFile = File(...)):
    content = await file.read()
    fname = file.filename or "upload"
    ext = Path(fname).suffix.lstrip(".").lower()
    if ext == "pdf":
        text, method = extract_pdf(content)
    elif ext == "docx":
        text = extract_docx(content)
        method = "docx"
    elif ext in ("png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"):
        text = ocr_image(Image.open(io.BytesIO(content)))
        method = "image_ocr"
    else:
        text = content.decode("utf-8", errors="ignore")
        method = "text"
    return {"success": True, "text": text, "length": len(text),
            "file_type": ext, "extraction_method": method,
            "engine_version": ENGINE_VERSION}


@app.post("/api/match-products")
async def match_products(req: MatchRequest):
    out = []
    for item in req.items:
        q = norm(item.get("product_name") or item.get("item")
                 or item.get("description"))
        ranked = []
        for p in req.company_products:
            names = [p.get("name", "")]
            if isinstance(p.get("aliases"), list):
                names.extend(p["aliases"])
            score = max(
                [100 if q == norm(n) else fuzz.WRatio(q, norm(n))
                 for n in names if n] or [0])
            ranked.append((score, p))
        ranked.sort(reverse=True, key=lambda x: x[0])
        best = ranked[0] if ranked else (0, None)
        second = ranked[1][0] if len(ranked) > 1 else 0
        accepted = bool(
            best[1] and best[0] >= 88
            and (best[0] - second >= 6 or best[0] >= 97))
        out.append({
            **item,
            "product_id": best[1].get("id") if accepted else None,
            "matched_product_name": best[1].get("name") if accepted else None,
            "match_confidence": round(best[0] / 100, 3),
            "match_status": "matched" if accepted else "review_required",
        })
    return {"success": True, "items": out,
            "matched_count": sum(1 for x in out
                                 if x["match_status"] == "matched"),
            "review_count": sum(1 for x in out
                                if x["match_status"] == "review_required")}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8000")))
